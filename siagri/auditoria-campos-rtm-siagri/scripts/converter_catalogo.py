"""Converte o catálogo XLSX do Siagri para TSV UTF-8 pesquisável.

Uso: python converter_catalogo.py origem.xlsx destino.tsv
Requer openpyxl. Não inclui o arquivo de origem no repositório.
"""

import csv
import sys
from pathlib import Path

from openpyxl import load_workbook


def texto(valor):
    if valor is None:
        return ""
    return " ".join(str(valor).split())


def converter(origem: Path, destino: Path):
    wb = load_workbook(origem, read_only=True, data_only=True)
    try:
        ws = wb.active
        destino.parent.mkdir(parents=True, exist_ok=True)
        with destino.open("w", encoding="utf-8", newline="") as arquivo:
            escritor = csv.writer(arquivo, delimiter="\t", lineterminator="\n")
            for linha in ws.iter_rows(values_only=True):
                escritor.writerow([texto(valor) for valor in linha[:7]])
    finally:
        wb.close()


if __name__ == "__main__":
    if len(sys.argv) != 3:
        raise SystemExit("Uso: python converter_catalogo.py origem.xlsx destino.tsv")
    converter(Path(sys.argv[1]), Path(sys.argv[2]))

